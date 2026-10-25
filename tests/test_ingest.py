import hashlib
import io
from types import SimpleNamespace
from unittest import mock

import docx
import pytest

from src.config import get_settings
from src.ingest import DocumentIngester, extract_text
from tests.conftest import EMBEDDING_DIM, make_bedrock_client


def minimal_pdf(text: str) -> bytes:
    """Build a one-page PDF with a text layer (no PDF library needed)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_ingester(existing_total_chunks: int = 0, chunk_size: int = 50):
    """DocumentIngester with mocked clients, bypassing the network checks in __init__."""
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    ingester = DocumentIngester.__new__(DocumentIngester)
    ingester.settings = get_settings()
    ingester.bedrock_client = make_bedrock_client()
    ingester.text_splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=0)
    ingester.pinecone_index = mock.Mock()

    def fetch(ids):
        vectors = {}
        if existing_total_chunks:
            vectors[ids[0]] = SimpleNamespace(metadata={"total_chunks": existing_total_chunks})
        return SimpleNamespace(vectors=vectors)

    ingester.pinecone_index.fetch.side_effect = fetch
    return ingester


def upserted_ids(index) -> list:
    return [vector[0] for call in index.upsert.call_args_list for vector in call.kwargs["vectors"]]


def deleted_ids(index) -> list:
    return [i for call in index.delete.call_args_list for i in call.kwargs["ids"]]


# --- text extraction ---------------------------------------------------------

def test_extract_text_plain_and_markdown():
    assert extract_text("notes.txt", "héllo".encode()) == "héllo"
    assert extract_text("README.MD", b"# Title") == "# Title"


def test_extract_text_docx_includes_tables():
    document = docx.Document()
    document.add_paragraph("Pump maintenance procedure")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "PT-101A"
    table.rows[0].cells[1].text = "Replace seal"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text("sop.docx", buffer.getvalue())

    assert "Pump maintenance procedure" in text
    assert "PT-101A | Replace seal" in text


def test_extract_text_pdf():
    assert "Torque spec 45 Nm" in extract_text("manual.pdf", minimal_pdf("Torque spec 45 Nm"))


def test_extract_text_rejects_unsupported_types():
    with pytest.raises(ValueError, match="Unsupported"):
        extract_text("sheet.xlsx", b"")


# --- chunk IDs and re-ingestion ----------------------------------------------

def test_chunk_ids_unchanged_so_existing_vectors_are_overwritten():
    ingester = make_ingester()
    expected = hashlib.md5(b"data/policy.txt:3").hexdigest()
    assert ingester.create_document_id("data/policy.txt", 3) == expected


def test_reingesting_shorter_document_deletes_stale_chunks():
    ingester = make_ingester(existing_total_chunks=5)
    text = "alpha " * 8 + "\n\n" + "beta " * 8  # two chunks at chunk_size=50

    count = ingester.ingest_document("data/policy.txt", "policy.txt", 100, text.encode())

    assert count == 2
    stale = [ingester.create_document_id("data/policy.txt", i) for i in (2, 3, 4)]
    assert deleted_ids(ingester.pinecone_index) == stale
    assert upserted_ids(ingester.pinecone_index) == [
        ingester.create_document_id("data/policy.txt", i) for i in (0, 1)
    ]


def test_first_ingest_deletes_nothing():
    ingester = make_ingester(existing_total_chunks=0)

    ingester.ingest_document("data/new.txt", "new.txt", 10, b"short text")

    ingester.pinecone_index.delete.assert_not_called()


def test_emptied_document_removes_all_its_chunks():
    ingester = make_ingester(existing_total_chunks=3)

    assert ingester.ingest_document("data/empty.txt", "empty.txt", 0, b"   ") == 0

    assert len(deleted_ids(ingester.pinecone_index)) == 3
    ingester.pinecone_index.upsert.assert_not_called()


def test_failed_chunk_embedding_fails_document_without_partial_upsert():
    ingester = make_ingester()
    calls = {"n": 0}

    def flaky_embedding(text):
        calls["n"] += 1
        if calls["n"] == 2:
            raise RuntimeError("throttled")
        return [0.1] * EMBEDDING_DIM

    ingester.generate_embedding = flaky_embedding
    text = "alpha " * 8 + "\n\n" + "beta " * 8 + "\n\n" + "gamma " * 8

    with pytest.raises(RuntimeError):
        ingester.ingest_document("data/doc.txt", "doc.txt", 100, text.encode())
    ingester.pinecone_index.upsert.assert_not_called()


def test_chunk_metadata():
    ingester = make_ingester()

    ingester.ingest_document("s3://bucket/a/doc.txt", "doc.txt", 42, b"some text")

    vector = ingester.pinecone_index.upsert.call_args.kwargs["vectors"][0]
    _, values, metadata = vector
    assert len(values) == EMBEDDING_DIM
    assert metadata["source"] == "s3://bucket/a/doc.txt"
    assert metadata["total_chunks"] == 1
    assert metadata["file_size"] == 42


# --- sources -------------------------------------------------------------------

def test_directory_ingestion_reports_failures(tmp_path):
    (tmp_path / "good.txt").write_text("good document")
    (tmp_path / "bad.txt").write_bytes(b"\xff\xfe not utf-8")
    (tmp_path / "ignored.csv").write_text("a,b")
    ingester = make_ingester()

    summary = ingester.ingest_directory(str(tmp_path))

    assert summary == {"files_processed": 1, "files_failed": 1, "chunks_ingested": 1}


def test_missing_directory_raises():
    with pytest.raises(FileNotFoundError):
        make_ingester().ingest_directory("/does/not/exist")


def test_s3_ingestion_reads_every_page():
    ingester = make_ingester()
    pages = [
        {"Contents": [{"Key": f"docs/{i}.txt", "Size": 5} for i in range(1000)]},
        {"Contents": [{"Key": "docs/last.txt", "Size": 5}, {"Key": "docs/image.png", "Size": 5}]},
    ]
    s3 = mock.Mock()
    s3.get_paginator.return_value.paginate.return_value = pages
    s3.get_object.side_effect = lambda Bucket, Key: {"Body": io.BytesIO(b"hello")}

    with mock.patch("src.ingest.boto3.client", return_value=s3):
        summary = ingester.ingest_s3_bucket("bucket", "docs/")

    assert summary["files_processed"] == 1001
    s3.get_paginator.assert_called_once_with("list_objects_v2")


# --- startup validation ----------------------------------------------------------

def test_embedding_dimension_mismatch_fails_fast():
    ingester = make_ingester()
    ingester.pinecone_index.describe_index_stats.return_value = SimpleNamespace(dimension=1536)

    with pytest.raises(ValueError, match="1024-dim vectors"):
        ingester._validate_embedding_dimension()


def test_matching_embedding_dimension_passes():
    ingester = make_ingester()
    ingester.pinecone_index.describe_index_stats.return_value = SimpleNamespace(dimension=EMBEDDING_DIM)

    ingester._validate_embedding_dimension()
