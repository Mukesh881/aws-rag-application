"""
Data ingestion script for the RAG pipeline.

This script processes documents from S3 or local storage, chunks them,
generates embeddings using AWS Bedrock, and stores them in Pinecone.
"""
import io
import json
import logging
import pathlib
import sys
from concurrent.futures import ThreadPoolExecutor
from typing import List, Dict, Any
from datetime import datetime, UTC
import hashlib

import boto3
from botocore.exceptions import ClientError, NoCredentialsError
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pinecone import Pinecone
from pypdf import PdfReader
import structlog

from .config import get_settings
from .utils import configure_logging, create_bedrock_client, serialize_index_stats

configure_logging()

logger = structlog.get_logger()

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}

# Pinecone accepts at most 1000 IDs per delete request
PINECONE_DELETE_BATCH = 1000


def extract_text(file_name: str, data: bytes) -> str:
    """
    Extract plain text from a document's raw bytes based on its extension.

    Args:
        file_name: File name or key, used to pick the parser
        data: Raw file content

    Returns:
        Extracted text (may be empty, e.g. for a scanned PDF without a text layer)
    """
    suffix = pathlib.PurePosixPath(file_name).suffix.lower()
    if suffix in (".txt", ".md"):
        return data.decode("utf-8")
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        document = docx.Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported file type: {suffix or file_name}")


class DocumentIngester:
    """Handles document ingestion into the RAG pipeline."""

    def __init__(self):
        """Initialize the ingester with AWS and Pinecone clients."""
        self.settings = get_settings()
        self._setup_clients()
        self._setup_text_splitter()

    def _setup_clients(self):
        """Initialize AWS Bedrock and Pinecone clients."""
        try:
            # Initialize Bedrock client
            self.bedrock_client = create_bedrock_client(self.settings.bedrock_region)
            logger.info("Bedrock client initialized", region=self.settings.bedrock_region)

            # Initialize Pinecone client
            self.pinecone_client = Pinecone(api_key=self.settings.pinecone_api_key)
            self.pinecone_index = self.pinecone_client.Index(self.settings.pinecone_index_name)
            logger.info("Pinecone client initialized", index=self.settings.pinecone_index_name)

            # Test Bedrock connection and check it matches the index
            self._validate_embedding_dimension()

        except NoCredentialsError:
            logger.error("AWS credentials not found. Please configure AWS CLI or set environment variables.")
            raise
        except Exception as e:
            logger.error("Failed to initialize clients", error=str(e))
            raise

    def _validate_embedding_dimension(self):
        """
        Embed a test string and check its size against the Pinecone index.

        A mismatch (e.g. Titan v1 at 1536 dims into a 1024-dim index) would
        otherwise only surface as upsert errors partway through ingestion.
        """
        try:
            embedding = self.generate_embedding("test")
            logger.info("Bedrock connection test successful", embedding_dim=len(embedding))
        except ClientError as e:
            if e.response['Error']['Code'] == 'AccessDeniedException':
                logger.error(
                    "Access denied to Bedrock. Ensure the model is enabled in your region",
                    model_id=self.settings.embed_model_id,
                    region=self.settings.bedrock_region
                )
            raise

        index_dimension = getattr(self.pinecone_index.describe_index_stats(), "dimension", None)
        if index_dimension and index_dimension != len(embedding):
            raise ValueError(
                f"Embedding model {self.settings.embed_model_id} produces {len(embedding)}-dim vectors "
                f"but Pinecone index {self.settings.pinecone_index_name} expects {index_dimension}"
            )

    def _setup_text_splitter(self):
        """Initialize the text splitter for chunking documents."""
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.settings.chunk_size,
            chunk_overlap=self.settings.chunk_overlap,
            length_function=len,
        )
        logger.info(
            "Text splitter configured",
            chunk_size=self.settings.chunk_size,
            chunk_overlap=self.settings.chunk_overlap
        )

    def generate_embedding(self, text: str) -> List[float]:
        """
        Generate embedding for a text using Bedrock Titan.

        Args:
            text: Input text to embed

        Returns:
            List of embedding values
        """
        try:
            response = self.bedrock_client.invoke_model(
                modelId=self.settings.embed_model_id,
                body=json.dumps({"inputText": text}),
                contentType="application/json",
                accept="application/json",
            )

            response_body = json.loads(response["body"].read())
            embedding = response_body["embedding"]

            logger.debug("Generated embedding", text_length=len(text), embedding_dim=len(embedding))
            return embedding

        except Exception as e:
            logger.error("Failed to generate embedding", error=str(e), text_length=len(text))
            raise

    def create_document_id(self, source: str, chunk_index: int) -> str:
        """
        Create a unique document ID for a chunk.

        Args:
            source: Source file path
            chunk_index: Index of the chunk within the document

        Returns:
            Unique document ID
        """
        content = f"{source}:{chunk_index}"
        return hashlib.md5(content.encode()).hexdigest()

    def build_chunks(self, source: str, file_name: str, file_size: int, text: str) -> List[Dict[str, Any]]:
        """
        Split text into chunks and embed them.

        Embeddings run in parallel (EMBED_CONCURRENCY); any chunk failure fails
        the whole document so it is never indexed with gaps.

        Args:
            source: Source path or S3 URI (determines chunk IDs)
            file_name: Base file name
            file_size: File size in bytes
            text: Document text

        Returns:
            List of processed chunks with embeddings and metadata
        """
        chunks = self.text_splitter.split_text(text)
        logger.info("Document split into chunks", source=source, chunk_count=len(chunks))

        with ThreadPoolExecutor(max_workers=self.settings.embed_concurrency) as pool:
            embeddings = list(pool.map(self.generate_embedding, chunks))

        ingested_at = datetime.now(UTC).isoformat()
        return [
            {
                "id": self.create_document_id(source, i),
                "values": embedding,
                "metadata": {
                    "text": chunk,
                    "source": source,
                    "chunk_index": i,
                    "total_chunks": len(chunks),
                    "file_name": file_name,
                    "file_size": file_size,
                    "ingested_at": ingested_at,
                    "text_length": len(chunk)
                }
            }
            for i, (chunk, embedding) in enumerate(zip(chunks, embeddings))
        ]

    def _existing_chunk_count(self, source: str) -> int:
        """Number of chunks currently indexed for a source, read from its first chunk."""
        first_id = self.create_document_id(source, 0)
        existing = self.pinecone_index.fetch(ids=[first_id]).vectors.get(first_id)
        if not existing or not existing.metadata:
            return 0
        return int(existing.metadata.get("total_chunks", 0))

    def _delete_stale_chunks(self, source: str, new_total: int, old_total: int):
        """
        Remove chunks left over from a previous, longer version of a document.

        Chunk IDs are deterministic per (source, chunk_index), so re-ingesting
        overwrites chunks 0..new_total-1; anything past that would otherwise
        keep serving outdated text.
        """
        stale_ids = [self.create_document_id(source, i) for i in range(new_total, old_total)]
        for i in range(0, len(stale_ids), PINECONE_DELETE_BATCH):
            self.pinecone_index.delete(ids=stale_ids[i:i + PINECONE_DELETE_BATCH])
        if stale_ids:
            logger.info("Deleted stale chunks", source=source, deleted=len(stale_ids))

    def ingest_document(self, source: str, file_name: str, file_size: int, data: bytes) -> int:
        """
        Extract, chunk, embed and upsert one document, replacing its previous version.

        Args:
            source: Source path or S3 URI (determines chunk IDs)
            file_name: Base file name, used to pick the parser
            file_size: File size in bytes
            data: Raw file content

        Returns:
            Number of chunks ingested
        """
        logger.info("Processing document", source=source)

        text = extract_text(file_name, data)
        if not text.strip():
            logger.warning("No text extracted from document", source=source)
        chunks = self.build_chunks(source, file_name, file_size, text) if text.strip() else []

        # Read the old chunk count before the upsert overwrites chunk 0
        old_total = self._existing_chunk_count(source)
        self.ingest_to_pinecone(chunks)
        self._delete_stale_chunks(source, len(chunks), old_total)

        logger.info("Document processing complete", source=source, processed_chunks=len(chunks))
        return len(chunks)

    def ingest_to_pinecone(self, chunks: List[Dict[str, Any]], batch_size: int = 100):
        """
        Ingest processed chunks into Pinecone.

        Args:
            chunks: List of processed chunks with embeddings
            batch_size: Number of chunks to upload in each batch
        """
        if not chunks:
            return

        try:
            # Process in batches
            for i in range(0, len(chunks), batch_size):
                batch = chunks[i:i + batch_size]

                # Prepare vectors for upsert
                vectors = [(chunk["id"], chunk["values"], chunk["metadata"]) for chunk in batch]

                # Upsert to Pinecone
                self.pinecone_index.upsert(vectors=vectors)

                logger.debug(
                    "Batch uploaded to Pinecone",
                    batch_start=i,
                    batch_size=len(batch),
                    total_uploaded=min(i + batch_size, len(chunks))
                )

        except Exception as e:
            logger.error("Failed to ingest to Pinecone", error=str(e))
            raise

    def _log_summary(self, message: str, summary: Dict[str, int]):
        """Log an ingestion summary together with the current index stats."""
        try:
            index_stats = serialize_index_stats(self.pinecone_index.describe_index_stats())
        except Exception as e:
            logger.warning("Failed to get index stats", error=str(e))
            index_stats = {"error": "Unable to retrieve stats"}
        logger.info(message, index_stats=index_stats, **summary)

    def ingest_directory(self, directory_path: str) -> Dict[str, int]:
        """
        Ingest all supported files from a directory.

        Args:
            directory_path: Path to the directory containing documents

        Returns:
            Summary with files_processed, files_failed and chunks_ingested
        """
        directory = pathlib.Path(directory_path)

        if not directory.is_dir():
            raise FileNotFoundError(f"Directory does not exist: {directory}")

        files = sorted(
            path for path in directory.rglob("*")
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
        )
        summary = {"files_processed": 0, "files_failed": 0, "chunks_ingested": 0}

        if not files:
            logger.warning(
                "No supported files found in directory",
                directory=str(directory),
                extensions=sorted(SUPPORTED_EXTENSIONS)
            )
            return summary

        logger.info("Starting directory ingestion", directory=str(directory), file_count=len(files))

        for file_path in files:
            try:
                summary["chunks_ingested"] += self.ingest_document(
                    source=str(file_path),
                    file_name=file_path.name,
                    file_size=file_path.stat().st_size,
                    data=file_path.read_bytes()
                )
                summary["files_processed"] += 1
            except Exception as e:
                summary["files_failed"] += 1
                logger.error("Failed to process file", file_path=str(file_path), error=str(e))

        self._log_summary("Ingestion complete", summary)
        return summary

    def ingest_s3_bucket(self, bucket_name: str, prefix: str = "") -> Dict[str, int]:
        """
        Ingest supported documents from an S3 bucket.

        Args:
            bucket_name: Name of the S3 bucket
            prefix: Optional prefix to filter objects

        Returns:
            Summary with files_processed, files_failed and chunks_ingested
        """
        s3_client = boto3.client('s3', region_name=self.settings.aws_region)

        # list_objects_v2 returns at most 1000 keys per call; paginate
        paginator = s3_client.get_paginator("list_objects_v2")
        objects = [
            obj
            for page in paginator.paginate(Bucket=bucket_name, Prefix=prefix)
            for obj in page.get("Contents", [])
            if pathlib.PurePosixPath(obj["Key"]).suffix.lower() in SUPPORTED_EXTENSIONS
        ]
        summary = {"files_processed": 0, "files_failed": 0, "chunks_ingested": 0}

        logger.info(
            "Found supported files in S3",
            bucket=bucket_name,
            prefix=prefix,
            file_count=len(objects)
        )
        if not objects:
            return summary

        for obj in objects:
            key = obj["Key"]
            try:
                data = s3_client.get_object(Bucket=bucket_name, Key=key)["Body"].read()
                summary["chunks_ingested"] += self.ingest_document(
                    source=f"s3://{bucket_name}/{key}",
                    file_name=key.split('/')[-1],
                    file_size=obj["Size"],
                    data=data
                )
                summary["files_processed"] += 1
            except Exception as e:
                summary["files_failed"] += 1
                logger.error("Failed to process S3 file", file_key=key, error=str(e))

        self._log_summary("S3 ingestion complete", summary)
        return summary


def main():
    """Main entry point for the ingestion script."""
    import argparse

    parser = argparse.ArgumentParser(description="Ingest documents into the RAG pipeline")
    parser.add_argument(
        "--source-type",
        choices=["local", "s3"],
        default="local",
        help="Source type for documents"
    )
    parser.add_argument(
        "--path",
        default="./data",
        help="Local directory path or S3 bucket name"
    )
    parser.add_argument(
        "--s3-prefix",
        default="",
        help="S3 prefix for filtering objects (only for S3 source)"
    )
    parser.add_argument(
        "--log-level",
        choices=["DEBUG", "INFO", "WARNING", "ERROR"],
        default="INFO",
        help="Logging level"
    )

    args = parser.parse_args()

    # Configure logging level
    logging.basicConfig(level=getattr(logging, args.log_level))

    try:
        ingester = DocumentIngester()

        if args.source_type == "local":
            summary = ingester.ingest_directory(args.path)
        else:
            summary = ingester.ingest_s3_bucket(args.path, args.s3_prefix)

    except Exception as e:
        logger.error("Ingestion failed", error=str(e))
        sys.exit(1)

    if summary["files_failed"]:
        logger.error("Ingestion finished with failures", **summary)
        sys.exit(1)
    logger.info("Ingestion completed successfully", **summary)


if __name__ == "__main__":
    main()
